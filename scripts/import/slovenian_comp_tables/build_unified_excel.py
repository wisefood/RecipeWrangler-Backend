"""Combine parsed Slovenian plant and meat tables into a review workbook."""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

REPO_ROOT = Path(__file__).resolve().parents[3]
DIR = REPO_ROOT / "data" / "Slovenian-Comp-Tables"

SECTION_EN = {
    "ENERGIJSKA VREDNOST": "Energy value",
    "UŽITNI DEL": "Edible portion",
    "GLAVNE SESTAVINE": "Main components",
    "ELEMENTI": "Elements",
    "ELEMENTI IN ELEMENTI V SLEDOVIH": "Elements and trace elements",
    "V VODI TOPNI VITAMINI": "Water-soluble vitamins",
    "V MAŠČOBAH TOPNI VITAMINI": "Fat-soluble vitamins",
    "VITAMINI": "Vitamins",
    "MONO IN DISAHARIDI": "Mono- and disaccharides",
    "MAŠČOBNE KISLINE": "Fatty acids",
    "STEROLI": "Sterols",
    "AMINOKISLINE": "Amino acids",
    "ORGANSKE KISLINE": "Organic acids",
    "OSTALE SESTAVINE": "Other components",
    "DRUGO": "Other",
}
SUBSECTION_EN = {
    "Makroelementi": "Macroelements",
    "Mikroelementi": "Microelements",
    "V maščobah topni": "Fat-soluble",
    "V vodi topni": "Water-soluble",
}
PLANT_CATEGORY_EN = {
    "ŽITA": "cereals and grain products",
    "IZDELKI": "milled cereal products",
    "KRUH": "bread",
    "ZELENJAVA": "vegetables",
    "SADJE": "fruits",
    "OLJNICE": "oilseeds and olive oil",
}
MEAT_CATEGORY_EN = {
    "SVINJINA": "pork", "GOVEDINA": "beef", "TELETINA": "veal",
    "OVČETINA": "mutton and lamb", "KONJSKO MESO": "horse meat",
    "KOZJE MESO": "goat meat", "KUNČJE MESO": "rabbit meat",
    "IZDELKI IZ MESA KLAVNIH ŽIVALI": "meat products",
    "PETELINJE MESO": "rooster meat", "PIŠČANČJE MESO": "chicken meat",
    "PURANJE MESO": "turkey meat", "GOSJE MESO": "goose meat",
    "RAČJE MESO": "duck meat", "IZDELKI IZ PERUTNINE": "poultry products",
    "JELENJAD": "deer (venison)", "SRNJAD": "roe deer",
    "DIVJI PRAŠIČ": "wild boar", "DIVJI ZAJEC": "hare", "PERJAD": "game birds",
    "MORSKE RIBE": "sea fish", "SLADKOVODNE RIBE": "freshwater fish",
    "RAKI IN MEHKUŽCI": "crustaceans and molluscs",
    "IZDELKI IZ RIB, RAKOV IN MEHKUŽCEV": "fish, crustacean and mollusc products",
}
MEAT_GROUP_EN = {
    "MESO KLAVNIH ŽIVALI": "Meat of slaughtered animals",
    "PERUTNINA": "Poultry",
    "DIVJAČINA": "Game",
    "RIBE, RAKI IN MEHKUŽCI": "Fish, crustaceans and molluscs",
}

crosswalk = {}
for row in csv.DictReader(open(f"{DIR}/nutrient_terms_crosswalk.csv", encoding="utf-8")):
    crosswalk[row["slovenian_term"]] = row["english_translation"]

def nutrient_en(term):
    if term == "Energijska vrednost, skupaj":
        return "Energy value, total"
    if term == "Užitni del":
        return "Edible portion"
    if term == "Sugars, total (computed)":
        return "Sugars, total (computed)"
    return crosswalk.get(term, term)

wb = Workbook()
ws = wb.active
ws.title = "Ingredients"

headers = [
    "book", "food_id", "food_name_en", "food_name_slo", "food_name_latin",
    "group_en", "category_en", "category_slo",
    "section_en", "section_slo", "subsection_en", "subsection_slo",
    "nutrient_en", "nutrient_slo", "unit", "avg", "min", "max",
]
ws.append(headers)
for cell in ws[1]:
    cell.font = Font(bold=True)

def add_book(book, foods_csv, nutrients_csv, category_map, group_map=None):
    foods = {r["food_id"]: r for r in csv.DictReader(open(f"{DIR}/{foods_csv}", encoding="utf-8"))}
    for r in csv.DictReader(open(f"{DIR}/{nutrients_csv}", encoding="utf-8")):
        fid = r["food_id"]
        f = foods[fid]
        cat_slo = f["category"]
        group_slo = f.get("group", "")
        section_slo = r["section"] or ""
        subsection_slo = r.get("subsection", "") or ""
        ws.append([
            book,
            fid,
            f["name_en"],
            f["name_slo"],
            f.get("name_latin", ""),
            (group_map.get(group_slo, group_slo) if group_map else ""),
            category_map.get(cat_slo, cat_slo),
            cat_slo,
            SECTION_EN.get(section_slo, section_slo),
            section_slo,
            SUBSECTION_EN.get(subsection_slo, subsection_slo),
            subsection_slo,
            nutrient_en(r["nutrient"]),
            r["nutrient"],
            r["unit"],
            r["avg"],
            r["min"],
            r["max"],
        ])

add_book("plant", "slovenian_foods.csv", "slovenian_nutrients_long.csv", PLANT_CATEGORY_EN)
add_book("meat", "slovenian_meat_foods.csv", "slovenian_meat_nutrients_long.csv", MEAT_CATEGORY_EN, MEAT_GROUP_EN)

widths = [8, 8, 32, 32, 26, 24, 24, 24, 26, 26, 18, 18, 34, 30, 8, 10, 10, 10]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

ws.freeze_panes = "A2"

out_path = f"{DIR}/slovenian_composition_tables_unified.xlsx"
wb.save(out_path)
print("wrote", out_path, "rows:", ws.max_row - 1)
